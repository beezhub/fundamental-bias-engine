"""Tests for the Reserve Bank of New Zealand provider, the seventh curve parser.

Issue #91 established two facts before any of this was written. The RBNZ
serves a JavaScript challenge to every network an unattended run on this
project can use, so the fixture was retrieved by the owner through a browser
and `tests/test_rbnz_b2_fixture.py` pins what it holds. And the 2-year lives in
a workbook whose four government bond columns differ only in the word before
"year", with the 1 and 5 year either side of the 2 year in the same unit and a
similar range, so a parser one column out returns a plausible wrong yield.

The tests here therefore check the parser against the pinned closes rather
than against itself, and check that the column is located by its series ID
rather than its position.

No test reaches the network. The workbook is served through `respx`, and the
cache lives under ``tmp_path``.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Iterator
from datetime import date
from pathlib import Path

import httpx
import openpyxl
import pytest
import respx

from fbe.config import DataConfig
from fbe.datasources import collect, registry
from fbe.datasources.base import SourceError
from fbe.datasources.curves import (
    PROVIDER_FETCHERS,
    PROVIDER_FOR_CURRENCY,
    RBNZ_B2_URL,
    RBNZ_DATA_SHEET,
    RBNZ_DROP_IN_FILENAME,
    RBNZ_SERIES_ID_ROW_LABEL,
    RBNZ_UNIT,
    RBNZ_UNIT_ROW_LABEL,
    TWO_YEAR_REFS,
    CurvesSource,
    RbnzSource,
)

FIXTURES = Path(__file__).parent / "fixtures"
WORKBOOK = (FIXTURES / "rbnz_hb2_daily_close.xlsx").read_bytes()

TWO_YEAR = "INM.DG102.NZZCF"
ONE_YEAR = "INM.DG101.NZZCF"
"""The 2-year and the column immediately to its left in the published file."""

START = date(2026, 9, 1)
END = date(2026, 9, 14)

PUBLISHED_CLOSES = {
    date(2026, 9, 8): 3.59,
    date(2026, 9, 9): 3.62,
    date(2026, 9, 10): 3.69,
    date(2026, 9, 11): 3.93,
    date(2026, 9, 14): 3.91,
}
"""Transcribed from `tests/test_rbnz_b2_fixture.py`, which read them off the
full workbook. Repeated here rather than imported so that this file asserts
against the publication and not against another test's constant."""

GAP_SESSION = date(2020, 5, 15)
"""A session the RBNZ published with the 2-year cell left blank."""

BLOCK_PAGE = (
    "<html><head><title>Website unavailable</title></head><body>"
    "Enable JavaScript and cookies to continue. Your access to the Reserve "
    "Bank website has been restricted.</body></html>"
)
"""What the RBNZ serves to a data-centre egress, at HTTP 403."""


@pytest.fixture
def source(tmp_path: Path) -> Iterator[CurvesSource]:
    # manual_dir under tmp_path as well, or the fetch reads whatever workbook
    # the owner has dropped into the repository's own data/manual (#283).
    curves = CurvesSource(
        DataConfig(cache_dir=tmp_path / "cache", manual_dir=tmp_path / "manual")
    )
    yield curves
    curves.close()


def _serve(content: bytes = WORKBOOK) -> None:
    respx.get(RBNZ_B2_URL).mock(return_value=httpx.Response(200, content=content))


def _rewritten(edit: object) -> bytes:
    """Return the fixture with one cell changed, for the tests that break it.

    Args:
        edit: A callable taking the ``Data`` worksheet.

    Returns:
        The edited workbook as bytes.

    """
    book = openpyxl.load_workbook(io.BytesIO(WORKBOOK))
    edit(book[RBNZ_DATA_SHEET])  # type: ignore[operator]
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# fetch_rbnz: the workbook
# ---------------------------------------------------------------------------


@respx.mock
def test_the_published_closes_are_reproduced(source: CurvesSource) -> None:
    """The five pinned sessions, to the cent the RBNZ published them at."""
    _serve()
    assert dict(source.fetch_rbnz(TWO_YEAR, START, END)) == PUBLISHED_CLOSES


@respx.mock
def test_the_documented_url_is_the_one_requested(source: CurvesSource) -> None:
    """The filename is ``hb2-daily-close.xlsx``. Two guessed paths were wrong."""
    route = respx.get(RBNZ_B2_URL).mock(
        return_value=httpx.Response(200, content=WORKBOOK)
    )
    source.fetch_rbnz(TWO_YEAR, START, END)
    assert route.called
    assert RBNZ_B2_URL.endswith("/b2/hb2-daily-close.xlsx")


@respx.mock
def test_the_column_is_found_by_series_id_not_by_position(
    source: CurvesSource,
) -> None:
    """Asking for the neighbouring tenor must return the neighbouring number."""
    _serve()
    two_year = dict(source.fetch_rbnz(TWO_YEAR, START, END))
    one_year = dict(source.fetch_rbnz(ONE_YEAR, START, END))
    assert two_year[date(2026, 9, 8)] == 3.59
    assert one_year[date(2026, 9, 8)] == 3.09
    assert two_year != one_year


@respx.mock
def test_a_reordered_column_is_still_found(source: CurvesSource) -> None:
    """Swap the 1 and 2 year columns and the 2 year must still come back.

    The RBNZ owns the column order and has 48 series to rearrange. A parser
    that remembered where the 2 year was last time would read the 1 year.
    """

    def swap(sheet: object) -> None:
        for row in sheet.iter_rows():  # type: ignore[attr-defined]
            row[3].value, row[4].value = row[4].value, row[3].value

    _serve(_rewritten(swap))
    assert dict(source.fetch_rbnz(TWO_YEAR, START, END)) == PUBLISHED_CLOSES


@respx.mock
def test_a_blank_session_is_absent_rather_than_zero(source: CurvesSource) -> None:
    """The 2020 gap. Read as 0.0 it hands MONETARY a zero front end for a year."""
    _serve()
    returned = dict(source.fetch_rbnz(TWO_YEAR, date(2020, 1, 1), date(2020, 12, 31)))
    assert GAP_SESSION not in returned
    assert 0.0 not in returned.values()


@respx.mock
def test_the_window_is_applied(source: CurvesSource) -> None:
    """The file is served whole. In a backtest ``end`` is the bar's as-of date."""
    _serve()
    returned = dict(source.fetch_rbnz(TWO_YEAR, date(2026, 9, 9), date(2026, 9, 10)))
    assert set(returned) == {date(2026, 9, 9), date(2026, 9, 10)}


@respx.mock
def test_rows_come_back_oldest_first(source: CurvesSource) -> None:
    _serve()
    days = [day for day, _value in source.fetch_rbnz(TWO_YEAR, START, END)]
    assert days == sorted(days)


@respx.mock
def test_an_unknown_series_id_raises_and_names_it(source: CurvesSource) -> None:
    _serve()
    with pytest.raises(SourceError) as excinfo:
        source.fetch_rbnz("INM.DG103.NZZCF", START, END)
    assert "INM.DG103.NZZCF" in str(excinfo.value)


@respx.mock
def test_a_workbook_without_the_series_id_row_raises(source: CurvesSource) -> None:
    """Without that row there is no way to tell one tenor's column from another."""

    def strip_label(sheet: object) -> None:
        for row in sheet.iter_rows():  # type: ignore[attr-defined]
            if row[0].value == RBNZ_SERIES_ID_ROW_LABEL:
                row[0].value = "Series"

    _serve(_rewritten(strip_label))
    with pytest.raises(SourceError) as excinfo:
        source.fetch_rbnz(TWO_YEAR, START, END)
    assert RBNZ_SERIES_ID_ROW_LABEL in str(excinfo.value)


@respx.mock
def test_a_workbook_without_the_data_sheet_raises(source: CurvesSource) -> None:
    book = openpyxl.load_workbook(io.BytesIO(WORKBOOK))
    book[RBNZ_DATA_SHEET].title = "Sheet1"
    out = io.BytesIO()
    book.save(out)
    _serve(out.getvalue())
    with pytest.raises(SourceError) as excinfo:
        source.fetch_rbnz(TWO_YEAR, START, END)
    assert RBNZ_DATA_SHEET in str(excinfo.value)


@respx.mock
def test_a_column_published_in_another_unit_is_refused(
    source: CurvesSource,
) -> None:
    """The registry promises percent. A yield in basis points is 100x and looks
    like a number a 2-year could take, so the unit row is checked, not assumed."""

    def rebase(sheet: object) -> None:
        for row in sheet.iter_rows():  # type: ignore[attr-defined]
            if row[0].value == RBNZ_UNIT_ROW_LABEL:
                row[4].value = "bp"

    _serve(_rewritten(rebase))
    with pytest.raises(SourceError) as excinfo:
        source.fetch_rbnz(TWO_YEAR, START, END)
    assert RBNZ_UNIT in str(excinfo.value)
    assert "bp" in str(excinfo.value)


@respx.mock
def test_the_block_page_raises_rather_than_being_read(
    source: CurvesSource, monkeypatch: pytest.MonkeyPatch
) -> None:
    """From a data-centre egress the RBNZ answers 403 with its own page.

    That is the answer every unattended run on this project gets, and it must
    be a named failure, not an empty series and not a cached body.
    """
    monkeypatch.setattr("time.sleep", lambda _seconds: None)
    respx.get(RBNZ_B2_URL).mock(return_value=httpx.Response(403, text=BLOCK_PAGE))
    with pytest.raises(SourceError):
        source.fetch_rbnz(TWO_YEAR, START, END)


@respx.mock
def test_a_body_that_is_not_a_workbook_raises(source: CurvesSource) -> None:
    """The block page at HTTP 200 would clear a status check and fail here."""
    respx.get(RBNZ_B2_URL).mock(return_value=httpx.Response(200, text=BLOCK_PAGE))
    with pytest.raises(SourceError):
        source.fetch_rbnz(TWO_YEAR, START, END)


# ---------------------------------------------------------------------------
# The registry and the dispatch tables
# ---------------------------------------------------------------------------


def test_the_nzd_two_year_is_now_a_fetchable_rbnz_ref() -> None:
    """Criterion 5 on the success branch: NZD leaves the manual route."""
    ref = registry.INDICATORS["yield_2y"].series["NZD"]
    assert ref.source == registry.SOURCE_RBNZ == "rbnz"
    assert ref.series_id == TWO_YEAR
    assert ref.unit == "percent"
    assert ref.fetchable
    assert ref.last_observed is not None


def test_the_derived_change_indicators_follow_the_level() -> None:
    """`_yield_change_series` is built from `YIELD_2Y`, so NZD moves with it."""
    for key in ("yield_2y_chg_1m", "yield_2y_chg_3m"):
        ref = registry.INDICATORS[key].series["NZD"]
        assert ref.source == registry.SOURCE_RBNZ
        assert ref.series_id == TWO_YEAR


def test_seven_of_eight_two_year_yields_are_fetchable() -> None:
    """The coverage figure the docs publish. CHF is the one that remains."""
    series = registry.INDICATORS["yield_2y"].series
    assert {code for code, ref in series.items() if ref.fetchable} == {
        "USD",
        "EUR",
        "GBP",
        "JPY",
        "CAD",
        "AUD",
        "NZD",
    }
    assert series["CHF"].source == registry.SOURCE_MANUAL


def test_rbnz_is_a_curve_source_with_a_fetcher() -> None:
    assert registry.SOURCE_RBNZ in registry.CURVE_SOURCES
    assert PROVIDER_FETCHERS[registry.SOURCE_RBNZ] == "fetch_rbnz"
    assert PROVIDER_FOR_CURRENCY["NZD"] == registry.SOURCE_RBNZ
    assert TWO_YEAR_REFS["NZD"] == (registry.SOURCE_RBNZ, TWO_YEAR)


@respx.mock
def test_fetch_emits_the_new_zealand_two_year_as_rbnz(
    source: CurvesSource,
) -> None:
    """Through the public entry point, with the provider named on the row."""
    _serve()
    emitted = source.fetch(["yield_2y"], ["NZD"], START, END)
    by_day = {o.period: o for o in emitted}
    assert by_day[date(2026, 9, 14)].value == 3.91
    assert {o.source for o in emitted} == {"rbnz"}
    assert {o.unit for o in emitted} == {"percent"}
    assert {o.currency for o in emitted} == {"NZD"}


@respx.mock
def test_provider_health_reports_the_rbnz(
    source: CurvesSource, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A run that cannot reach the RBNZ sees ``None`` here, and that is the
    diagnosis: the block, not a parsing fault."""
    monkeypatch.setattr("time.sleep", lambda _seconds: None)
    _serve()
    respx.route().mock(return_value=httpx.Response(503, text="down"))
    health = source.provider_health()
    assert health["rbnz"] == date(2026, 9, 14)


def test_the_docs_record_the_vantage_and_the_verdict() -> None:
    """The recording criteria of #91, checked as text so they cannot go stale
    without this failing."""
    text = (Path(__file__).parents[1] / "docs" / "data-sources.md").read_text(
        encoding="utf-8"
    )
    assert "hb2-daily-close.xlsx" in text
    assert "residential" in text
    assert "| `yield_2y` | 7/8 |" in text
    assert "worth one attempt from another network" not in text


# ---------------------------------------------------------------------------
# The drop-in file: the owner's browser download, read before the wire (#283)
# ---------------------------------------------------------------------------


@pytest.fixture
def rbnz(tmp_path: Path) -> Iterator[RbnzSource]:
    """The provider source, with a manual directory the tests can drop into."""
    (tmp_path / "manual").mkdir()
    provider = RbnzSource(
        DataConfig(cache_dir=tmp_path / "cache", manual_dir=tmp_path / "manual")
    )
    yield provider
    provider.close()


def _drop_in(tmp_path: Path, content: bytes = WORKBOOK) -> Path:
    path = tmp_path / "manual" / RBNZ_DROP_IN_FILENAME
    path.write_bytes(content)
    return path


@respx.mock
def test_the_drop_in_file_is_read_and_no_request_is_made(
    rbnz: RbnzSource, tmp_path: Path
) -> None:
    """With the file in place the RBNZ's refusal is never even seen.

    The block keys on the TLS handshake, so no request the engine can make
    will pass; the file is the route, and the wire is not tried while it
    exists. Every route answers 403 here to prove no request goes out.
    """
    _drop_in(tmp_path)
    route = respx.route().mock(return_value=httpx.Response(403, text=BLOCK_PAGE))

    emitted = rbnz.fetch(["yield_2y"], ["NZD"], START, END)

    assert route.call_count == 0
    assert {o.period: o.value for o in emitted} == PUBLISHED_CLOSES
    assert {o.source for o in emitted} == {"rbnz"}


@respx.mock
def test_without_the_file_the_wire_is_fetched_as_before(rbnz: RbnzSource) -> None:
    route = respx.get(RBNZ_B2_URL).mock(
        return_value=httpx.Response(200, content=WORKBOOK)
    )

    emitted = rbnz.fetch(["yield_2y"], ["NZD"], START, END)

    assert route.call_count == 1
    assert {o.period: o.value for o in emitted} == PUBLISHED_CLOSES


@respx.mock
def test_the_drop_in_file_is_served_offline(tmp_path: Path) -> None:
    """The file is not the cache. An offline run with a cold cache reads it."""
    (tmp_path / "manual").mkdir()
    path = _drop_in(tmp_path)
    route = respx.route().mock(return_value=httpx.Response(403, text=BLOCK_PAGE))
    offline = RbnzSource(
        DataConfig(
            cache_dir=tmp_path / "cache", manual_dir=tmp_path / "manual", offline=True
        )
    )
    try:
        emitted = offline.fetch(["yield_2y"], ["NZD"], START, END)
    finally:
        offline.close()

    assert path.exists()
    assert route.call_count == 0
    assert {o.period: o.value for o in emitted} == PUBLISHED_CLOSES


@respx.mock
def test_a_stale_drop_in_file_is_refused_by_name(
    rbnz: RbnzSource, tmp_path: Path
) -> None:
    """A download nobody refreshed must be a loud absence, not a quiet yield.

    The fixture's newest session is 2026-09-14. As of 2026-12-01 that is 78
    days old against a 9-day allowance, which is the age at which the leg
    would carry no weight anyway; refusing it there means the operator is
    told to download a fresh copy rather than left reading a coverage figure.
    """
    path = _drop_in(tmp_path)
    respx.route().mock(return_value=httpx.Response(403, text=BLOCK_PAGE))

    with pytest.raises(SourceError) as excinfo:
        rbnz.fetch(["yield_2y"], ["NZD"], date(2026, 11, 1), date(2026, 12, 1))

    message = str(excinfo.value)
    assert "stale" in message
    assert str(path) in message
    assert "2026-09-14" in message
    assert "9 days" in message
    assert RBNZ_B2_URL in message


@respx.mock
def test_a_backtest_window_before_the_newest_session_is_not_stale(
    rbnz: RbnzSource, tmp_path: Path
) -> None:
    """Freshness is the file's newest session against ``end``, not the window's.

    A current file asked for 2026-09-01 to 2026-09-10 is not stale: its newest
    session is after ``end``. Only a file whose newest session is too far
    before ``end`` is refused.
    """
    _drop_in(tmp_path)
    respx.route().mock(return_value=httpx.Response(403, text=BLOCK_PAGE))

    emitted = rbnz.fetch(["yield_2y"], ["NZD"], date(2026, 9, 1), date(2026, 9, 10))

    assert {o.period for o in emitted} == {
        d for d in PUBLISHED_CLOSES if d <= date(2026, 9, 10)
    }


@respx.mock
def test_a_stale_drop_in_file_is_a_failed_source_through_the_collector(
    tmp_path: Path,
) -> None:
    """The refresh line reads ``rbnz failed (...)`` and NZD lands nowhere."""
    (tmp_path / "manual").mkdir()
    _drop_in(tmp_path)
    respx.route().mock(return_value=httpx.Response(403, text=BLOCK_PAGE))

    result = collect.collect(
        DataConfig(cache_dir=tmp_path / "cache", manual_dir=tmp_path / "manual"),
        start=date(2026, 11, 1),
        end=date(2026, 12, 1),
        sources=(RbnzSource,),
        indicators=["yield_2y"],
        currencies=["NZD"],
    )

    outcome = {o.source: o for o in result.outcomes}["rbnz"]
    assert outcome.status is collect.SourceStatus.FAILED
    assert "stale" in outcome.detail
    assert RBNZ_DROP_IN_FILENAME in outcome.detail
    assert result.observations == ()


@respx.mock
def test_a_drop_in_file_that_is_not_table_b2_is_refused_by_the_parser(
    rbnz: RbnzSource, tmp_path: Path
) -> None:
    """The same checks that refuse a wrong body refuse a wrong file.

    The series ID is looked up by name, so a workbook without it is named as
    such rather than read by column position.
    """
    _drop_in(tmp_path)
    respx.route().mock(return_value=httpx.Response(403, text=BLOCK_PAGE))

    with pytest.raises(SourceError, match="does not carry NOT.A.SERIES"):
        rbnz.fetch_rbnz("NOT.A.SERIES", START, END)


@respx.mock
def test_a_drop_in_file_that_is_not_a_workbook_is_refused_not_fallen_back(
    rbnz: RbnzSource, tmp_path: Path
) -> None:
    """A broken file is a broken file, not a reason to try the wire.

    Falling back would hide the broken file behind a 403 the operator has
    already learned to ignore.
    """
    _drop_in(tmp_path, content=b"not a workbook")
    route = respx.route().mock(return_value=httpx.Response(200, content=WORKBOOK))

    with pytest.raises(SourceError, match="not a workbook"):
        rbnz.fetch(["yield_2y"], ["NZD"], START, END)
    assert route.call_count == 0


def _with_single_cell_dimension(content: bytes) -> bytes:
    """Return the workbook with every sheet's dimension declared as ``A1``.

    That is how the RBNZ actually publishes the file. The fixture was re-saved
    by openpyxl, which wrote the true dimension, so nothing before #283 read a
    file shaped like the real download.
    """
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(content)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            body = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                body = re.sub(
                    rb"<dimension ref=\"[^\"]*\"/>", b'<dimension ref="A1"/>', body
                )
            target.writestr(item, body)
    return out.getvalue()


@respx.mock
def test_the_published_file_declares_one_cell_and_is_still_read(
    rbnz: RbnzSource, tmp_path: Path
) -> None:
    """The real download says each sheet is one cell; the parser must not trust it.

    Found on the first real drop-in, 2026-09-24: 2190 rows and 48 columns
    behind a ``<dimension ref="A1"/>``, and the read-only reader yielded one
    empty row, so the file was refused as carrying no series-ID row.
    """
    shaped_like_the_download = _with_single_cell_dimension(WORKBOOK)
    assert shaped_like_the_download != WORKBOOK
    _drop_in(tmp_path, content=shaped_like_the_download)
    respx.route().mock(return_value=httpx.Response(403, text=BLOCK_PAGE))

    emitted = rbnz.fetch(["yield_2y"], ["NZD"], START, END)

    assert {o.period: o.value for o in emitted} == PUBLISHED_CLOSES
